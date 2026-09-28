import pandas as pd
import openpyxl
import os

files = ['excel1.xlsx', 'excel2.xlsx', 'excel3.xlsx', 'excel4.xlsx']

for file in files:
    if not os.path.exists(file):
        print(f"File {file} not found")
        continue
    
    print(f"==================================================")
    print(f"ANALYZING: {file}")
    print(f"==================================================")
    
    wb = openpyxl.load_workbook(file, data_only=False)
    for sheet_name in wb.sheetnames:
        print(f"\n--- SHEET: {sheet_name} ---")
        ws = wb[sheet_name]
        
        # Get headers (assume row 1 or 2, let's just get the first 5 rows to see structure)
        data = []
        for row in ws.iter_rows(min_row=1, max_row=5, values_only=False):
            row_data = []
            for cell in row:
                if cell.data_type == 'f':
                    row_data.append(f"FORMULA: {cell.value}")
                else:
                    row_data.append(str(cell.value))
            data.append(row_data)
            
        for i, row in enumerate(data):
            print(f"Row {i+1}: {row}")
            
        print("\nMerged Cells:", ws.merged_cells.ranges)